"""Output: results.xlsx (colour-coded) and results.jsonl. Never touches the input file."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .aspects import nlp_status

HELPER_COLS = {"_stances", "like_none", "dislike_none"}
WIDE = {"explanation": 70, "like_evidence": 45, "dislike_evidence": 45, "rule_ids": 28, "extra_flags": 24}


def review_ready(out: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    skip = set(cfg["colors"]["not_a_problem"])
    order = {lab: i for i, lab in enumerate(cfg["weights"]["severity_order"])}
    r = out[~out["label"].isin(skip)].copy()
    r["_o"] = r["label"].map(order)
    return r.sort_values(["_o", "confidence"], ascending=[True, False]).drop(columns="_o")


def build_sheets(out: pd.DataFrame, cfg: dict) -> dict[str, pd.DataFrame]:
    results = out.drop(columns=[c for c in HELPER_COLS if c in out.columns])
    review = review_ready(results, cfg)
    review["human_label"] = ""
    order = cfg["weights"]["severity_order"]
    counts = out["label"].value_counts().reindex(order).dropna().astype(int).rename_axis("label").reset_index(name="rows")
    counts["meaning"] = counts["label"].map(lambda l: cfg["colors"]["labels"].get(l, {}).get("meaning", ""))
    bucket = out.groupby(["label", "bucket"]).size().unstack(fill_value=0).reset_index()
    none_share = round(100 * (out["dislike_none"].isin(["NONE", "HEDGED_NONE"])).mean(), 1) if "dislike_none" in out else None
    notes = pd.DataFrame({"item": ["% Dislike answers none-ish", "spaCy / parser status", "corpus-generic nouns ignored"],
                          "value": [none_share, nlp_status() or "OK", ", ".join(out.attrs.get("generic_nouns", [])) or "-"]})
    asp = out[out["aspect"].astype(str) != ""]["aspect"].value_counts().rename_axis("aspect").reset_index(name="times_matched")
    return {"Results": results, "Review": review, "Summary": counts, "Summary_by_bucket": bucket,
            "Summary_notes": notes, "Aspects": asp}


def _style(ws, frame: pd.DataFrame, cfg: dict) -> None:
    colors = cfg["colors"]["labels"]
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="404040")
    for i, name in enumerate(frame.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDE.get(name, 45 if frame[name].astype(str).str.len().max() > 40 else 16)
    if "label" in frame.columns:
        li = list(frame.columns).index("label") + 1
        for r in range(2, len(frame) + 2):
            lab = ws.cell(r, li).value
            spec = colors.get(lab)
            if spec:
                for c in range(1, len(frame.columns) + 1):
                    cell = ws.cell(r, c)
                    cell.fill = PatternFill("solid", fgColor=spec["fill"])
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(r, li).font = Font(bold=True, color=spec["strong"])
    ws.auto_filter.ref = ws.dimensions


def _legend(ws, cfg: dict) -> None:
    ws.append(["label", "colour", "meaning"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="404040")
    for lab, spec in cfg["colors"]["labels"].items():
        ws.append([lab, "", spec["meaning"]])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=spec["fill"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color=spec["strong"])
    for col, w in zip("ABC", (24, 10, 48)):
        ws.column_dimensions[col].width = w


def write_workbook(sheets: dict[str, pd.DataFrame], cfg: dict, target) -> None:
    """target: path or file-like. Writes all sheets plus a colour Legend, styled."""
    with pd.ExcelWriter(target, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name[:31], index=False)
            _style(xw.sheets[name[:31]], frame, cfg)
        _legend(xw.book.create_sheet("Legend"), cfg)


def write_report(out: pd.DataFrame, cfg: dict, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(build_sheets(out, cfg), cfg, path)
    with open(path.with_suffix(".jsonl"), "w", encoding="utf-8") as fh:
        for idx, row in out.iterrows():
            fh.write(json.dumps({"row": int(idx) if isinstance(idx, int) else str(idx), "label": row["label"],
                                 "confidence": row["confidence"], "stances": row.get("_stances", [])}, default=str) + "\n")
    return path
